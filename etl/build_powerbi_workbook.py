"""
Build a lean, Power BI-ready Excel workbook from the SQL analysis query outputs.
Each sheet = one native Excel Table, sourced directly from a documented query
in sql/03_analysis_queries.sql. Small enough for browser upload, no gateway needed.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, PatternFill
import os

SRC = "data/powerbi"
OUT = "data/MSHA_PowerBI_Data.xlsx"
os.makedirs(os.path.dirname(OUT), exist_ok=True)

SHEETS = [
    ("trend_national", "trend_national_yearly.csv"),
    ("trend_coal_metal", "trend_coal_vs_metal.csv"),
    ("by_state", "by_state.csv"),
    ("top_classification", "top_classification.csv"),
    ("top_body_part", "top_body_part.csv"),
    ("experience_bands", "experience_bands.csv"),
    ("riskiest_mines", "riskiest_mines.csv"),
]

wb = Workbook()
wb.remove(wb.active)

for sheet_name, fname in SHEETS:
    df = pd.read_csv(os.path.join(SRC, fname))
    ws = wb.create_sheet(title=sheet_name)

    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2A78D6")

    for row in df.itertuples(index=False, name=None):
        ws.append(row)

    n_rows, n_cols = len(df) + 1, len(df.columns)
    ref = f"A1:{get_column_letter(n_cols)}{n_rows}"
    tbl = Table(displayName=f"tbl_{sheet_name}", ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tbl)

    for j, col in enumerate(df.columns, start=1):
        ws.column_dimensions[get_column_letter(j)].width = min(max(len(str(col)) + 2, 12), 32)

    print(f"{sheet_name}: {n_rows-1} rows -> tbl_{sheet_name}")

wb.save(OUT)
print("Saved:", OUT, "-", round(os.path.getsize(OUT)/1024, 1), "KB")
